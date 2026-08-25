"""Đọc tài liệu thô (PDF/DOCX/XLSX/TXT) và trả về text thuần cho agent xử lý.

Phase 1: chỉ hỗ trợ file text-based. Ảnh chụp lịch trình (OCR) dời sang Phase 2
— xem README mục "Trạng thái hiện tại".
"""

from pathlib import Path

import openpyxl
from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt"}


class UnsupportedFileTypeError(ValueError):
    pass


def extract_text(file_path: str | Path) -> str:
    """Trích xuất text thuần từ 1 file. Ném UnsupportedFileTypeError nếu định
    dạng chưa hỗ trợ ở Phase 1 (vd: .doc cũ, ảnh .jpg/.png cần OCR)."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _extract_pdf(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".xlsx":
        return _extract_xlsx(path)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")

    raise UnsupportedFileTypeError(
        f"Định dạng '{ext}' chưa hỗ trợ ở Phase 1 (hỗ trợ: {', '.join(sorted(SUPPORTED_EXTENSIONS))})"
    )


def _extract_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages_text).strip()


def _extract_docx(path: Path) -> str:
    document = Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]

    # Nhiều lịch trình tour trình bày dạng bảng (giờ | hoạt động | ghi chú) —
    # đọc luôn nội dung bảng, không chỉ đoạn văn.
    for table in document.tables:
        for row in table.rows:
            cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells_text:
                parts.append(" | ".join(cells_text))

    return "\n".join(parts).strip()


def _extract_xlsx(path: Path) -> str:
    """Đọc toàn bộ sheet dạng bảng text — phù hợp cho danh sách khách
    (tên, SĐT, số ghế, số phòng, ghi chú ăn uống theo từng cột)."""
    workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    lines: list[str] = []

    for sheet in workbook.worksheets:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                lines.append(" | ".join(cells))

    return "\n".join(lines).strip()
